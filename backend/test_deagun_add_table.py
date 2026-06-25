from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select
import time
from openpyxl import Workbook
import sys
sys.stdout.reconfigure(encoding='utf-8')
import os
from html_reporter import generate_html_report

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'doc', 'datatest')))
from data_deagun_add_table import TEST_CASES_ADD_TABLE

def login_admin(driver, wait):
    driver.get("https://deagun.dpdns.org/")
    time.sleep(2)
    try:
        nav_login = wait.until(EC.element_to_be_clickable((By.ID, "navLoginBtn")))
        nav_login.click()
        time.sleep(1)
        wait.until(EC.presence_of_element_located((By.ID, "loginUsername"))).send_keys("admin")
        driver.find_element(By.ID, "loginPassword").send_keys("admin123")
        driver.find_element(By.XPATH, "//form[@id='loginForm']//button[@type='submit']").click()
        time.sleep(3)
    except:
        pass

def run_add_table_test(driver, wait, data):
    print(f"\n==================================================")
    print(f"BẮT ĐẦU CHẠY: [{data['id']}] - {data['name']}")
    
    test_result = "Fail"
    is_success = False
    
    try:
        driver.get("https://deagun.dpdns.org/tables.html")
        time.sleep(2)
        
        open_modal_btn = wait.until(EC.element_to_be_clickable((By.XPATH, "//button[@data-bs-target='#addTableModal']")))
        open_modal_btn.click()
        time.sleep(1)
        
        name_input = wait.until(EC.presence_of_element_located((By.ID, "tableName")))
        name_input.clear()
        if data['table_name']:
            name_input.send_keys(data['table_name'])
            
        area_select = Select(driver.find_element(By.ID, "tableLocation"))
        if data['area']:
            try:
                area_select.select_by_visible_text(data['area'])
            except:
                area_select.select_by_index(1) # fallback
            
        capacity_input = driver.find_element(By.ID, "tableCapacity")
        capacity_input.clear()
        if data['capacity']:
            capacity_input.send_keys(data['capacity'])
        
        save_btn = driver.find_element(By.XPATH, "//form[@id='addTableForm']//button[@type='submit']")
        save_btn.click()
        
        time.sleep(2)
        
        # In modern web, validation prevents submission, or there's a toast
        if not save_btn.is_displayed():
            is_success = True
        else:
            try:
                # Check for toast success
                toast = driver.find_element(By.ID, "toastMessage")
                if toast.is_displayed():
                    is_success = True
            except:
                pass
            
            # Check for HTML validation error
            try:
                invalid_elements = driver.find_elements(By.CSS_SELECTOR, "input:invalid, select:invalid")
                if len(invalid_elements) > 0:
                    is_success = False
            except:
                pass

        if is_success == data['expect_success']:
            test_result = "Pass"
            print(f"=> RESULT: PASS")
        else:
            test_result = "Fail"
            print(f"=> RESULT: FAIL")
            
    except Exception as e:
        print(f"=> LỖI THỰC THI: {str(e)[:50]}")
            
    return test_result

def main():
    driver = webdriver.Chrome()
    driver.implicitly_wait(5)
    wait = WebDriverWait(driver, 10)
    
    login_admin(driver, wait)
    results = []
    
    for case in TEST_CASES_ADD_TABLE:
        status = run_add_table_test(driver, wait, case)
        results.append({
            "ID": case['id'],
            "Testcase": case['name'],
            "Pass/Fail": status
        })
        
    wb = Workbook()
    ws = wb.active
    ws.title = "DEAGUN Add Table Report"
    ws.append(["ID", "Testcase", "Pass/Fail"])
    for r in results:
        ws.append([r['ID'], r['Testcase'], r['Pass/Fail']])
    
    report_file = "report_deagun_add_table.xlsx"
    wb.save(report_file)
    
    html_report_file = report_file.replace('.xlsx', '.html')
    generate_html_report("DEAGUN - Thêm bàn", results, html_report_file)
    
    print(f"\n==================================================")
    print(f"Hoàn thành! Báo cáo đã lưu: {report_file} và {html_report_file}")
    driver.quit()

if __name__ == "__main__":
    main()
